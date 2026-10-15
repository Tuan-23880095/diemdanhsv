#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Phân tích dữ liệu nghiên cứu đi kèm Workshop “Làm chủ thời gian” — Tân sinh viên K2026, HCMUS.

Đầu vào : tệp CSV tải từ bảng diễn giả (nút “Tải dữ liệu nghiên cứu (CSV)”), mỗi dòng = một mã ẩn danh.
Đầu ra  : một thư mục gồm
            bao_cao_phan_tich.docx   — báo cáo tiếng Việt, bảng theo kiểu APA, câu kết quả viết sẵn
            ket_qua_phan_tich.xlsx   — mọi bảng số + dữ liệu đã chấm điểm + câu trả lời mở
            du_lieu_da_cham.csv      — dữ liệu đã đảo chiều/tính điểm thang (dùng cho SPSS, R, AMOS)
            cfa_tmq.R                — mô hình CFA 3 nhân tố cho TMQ (chạy bằng R + lavaan)
            bieu_do/*.png            — biểu đồ

Cách dùng (Windows, cần Python 3.10+ với pandas, scipy, matplotlib, python-docx, openpyxl):
    python phan_tich.py du_lieu_nghien_cuu.csv
    python phan_tich.py du_lieu_nghien_cuu.csv --ra ket_qua --thap-phan .
    python phan_tich.py --gia-lap GIA_LAP_du_lieu.csv        # tạo dữ liệu giả lập để chạy thử

Khóa chấm TMQ: Britton & Tesser (1991), Bảng 1 — câu 8, 10, 11, 13, 16 chấm đảo (6 − x).
Dữ liệu giả lập dùng mã bắt đầu bằng “ZZ” (không có trong tên tiếng Việt) và mọi đầu ra
từ dữ liệu đó đều được đóng dấu “DỮ LIỆU GIẢ LẬP”.
"""
from __future__ import annotations

import argparse
import datetime as dt
import math
import os
import re
import sys
import warnings

import numpy as np
import pandas as pd
from scipy import stats

warnings.filterwarnings('ignore', category=RuntimeWarning)

# ----------------------------------------------------------------------------- công cụ đo
TMQ_SUB = ['SRP'] * 7 + ['TA'] * 6 + ['LRP'] * 5
TMQ_REV = [8, 10, 11, 13, 16]
TMQ_NAME = {'SRP': 'Lập kế hoạch ngắn hạn', 'TA': 'Thái độ với thời gian', 'LRP': 'Lập kế hoạch dài hạn', 'TOT': 'TMQ toàn thang'}
KT_KEY = [2, 1, 2, 0, 1, 1, 2, 2]
KT_NAME = ['K1. Mục tiêu SMART', 'K2. Ô của việc đọc trước bài', 'K3. Xử lý việc khẩn cấp-không quan trọng',
           'K4. Phiên Pomodoro', 'K5. Vùng đệm trong lịch tuần', 'K6. “Ăn con ếch”', 'K7. 66 ngày (Lally, 2010)',
           'K8. 50 giờ/tín chỉ']
SE_NAME = ['Viết mục tiêu cụ thể, đo được', 'Phân biệt quan trọng/khẩn cấp', 'Lập và theo kế hoạch tuần',
           'Từ chối lời nhờ vả lệch kế hoạch', 'Bắt đầu ngay việc khó']
RE_NAME = ['Nội dung hữu ích', 'Phù hợp tân sinh viên', 'Hoạt động trên điện thoại giúp hiểu bài', 'Sẽ giới thiệu cho bạn bè']
YD_NAME = ['Áp dụng ≥ 1 kỹ thuật trong 2 tuần', 'Lập kế hoạch cho tuần tới', 'Tin rằng kỹ thuật giúp học tốt hơn']
AD_NAME = ['Mục tiêu SMART', 'Ma trận Eisenhower', 'Lịch tuần, khối giờ tự học', 'Rà soát tuần', 'Pomodoro',
           'Quy tắc 2 phút', '“Ăn con ếch”']
AD_SCALE = ['Chưa lần nào', '1–2 lần', 'Vài lần mỗi tuần', 'Gần như mỗi ngày']
TR_NAME = ['Quên mất', 'Không đủ thời gian', 'Lịch học/làm thêm thay đổi', 'Thiếu động lực', 'Kỹ thuật không hợp',
           'Không biết bắt đầu từ đâu', 'Bị cuốn vào điện thoại, MXH', 'Khác']
THAYDOI = ['Kém hơn nhiều', 'Kém hơn một chút', 'Không thay đổi', 'Tốt hơn một chút', 'Tốt hơn nhiều']
DEMO = {
    'gioi': ('Giới tính', ['Nam', 'Nữ', 'Khác', 'Không muốn trả lời']),
    'nganh': ('Nhóm ngành', ['Toán – Thống kê – KHDL', 'CNTT – Trí tuệ nhân tạo', 'Vật lý – Điện tử – Hạt nhân – Bán dẫn',
                             'Hóa học – CN hóa học', 'Sinh học – CN sinh học', 'KH & CN vật liệu',
                             'Địa chất – Hải dương – Môi trường', 'Công nghệ giáo dục', 'Khác', 'Không muốn trả lời']),
    'noio': ('Nơi ở', ['Cùng gia đình', 'Ký túc xá', 'Nhà trọ', 'Nhà người thân', 'Khác']),
    'lamthem': ('Làm thêm mỗi tuần', ['Không làm thêm', 'Dưới 10 giờ', '10–20 giờ', 'Trên 20 giờ']),
    'tinchi': ('Số tín chỉ học kỳ này', ['Dưới 15', '15–18', '19–22', 'Trên 22', 'Không rõ']),
    'coso': ('Cơ sở học chính', ['Linh Trung', 'Nguyễn Văn Cừ', 'Cả hai']),
    'dahoc': ('Đã từng học về QLTG', ['Chưa', 'Rồi']),
}
SHAPES = {'vuong': 'Vuông', 'tamgiac': 'Tam giác', 'chunhat': 'Chữ nhật', 'tron': 'Tròn', 'luonsong': 'Lượn sóng'}
EISEN_KEY = {'c1': 1, 'c2': 2, 'c3': 3, 'c4': 4, 'c5': 2, 'c6': 3, 'c7': 1, 'c8': 2}
EISEN_NAME = {'c1': 'Bài nộp Moodle 23:59 tối nay', 'c2': 'Đọc trước chương 2 Giải tích', 'c3': 'Lấy hộ đơn hàng',
              'c4': 'Lướt Reels “thêm 5 phút”', 'c5': 'Tập thể dục 30 phút', 'c6': 'Group chat 40 tin',
              'c7': 'Mai thi, chưa ôn', 'c8': 'Gặp cố vấn học tập'}
WHEEL = ['Học tập', 'Sức khỏe & giấc ngủ', 'Tài chính', 'Gia đình', 'Bạn bè & quan hệ', 'Phát triển bản thân',
         'Giải trí & nghỉ ngơi', 'Không gian sống']
PLEDGES = ['Ngủ đủ 7 tiếng', '3 việc quan trọng cho ngày mai', '“Ăn con ếch” trước 10 giờ', 'Tắt thông báo MXH khi học',
           '2 phiên Pomodoro/ngày', 'Rà soát tuần tối Chủ nhật', 'Xong bài trước hạn 24 giờ', 'Nói “không” với lời mời lệch mục tiêu']
CAM = ['Dư dả', 'Vừa đủ', 'Hơi đuối', 'Ngập đầu']
MAN = ['Dưới 2 giờ', '2–4 giờ', '4–6 giờ', '6–8 giờ', 'Trên 8 giờ']

MAU = {'luc': '#0F6B55', 'lam': '#1F4E79', 'vang': '#B7811A', 'do': '#D2492A', 'xam': '#5B6B7C', 'muc': '#10212B'}


def csv_header() -> list[str]:
    """Danh sách cột đúng như ws_nc_csv() trong workshop/nc-lib.php."""
    h = ['ma_an_danh', 'so_thiet_bi', 'co_t0', 'co_t1', 'co_t2', 'co_t3']
    h += ['t0_' + k for k in DEMO] + ['t0_thoi_gian']
    tmq = lambda w: [f'{w}_tmq{i:02d}_{TMQ_SUB[i - 1]}' + ('_R' if i in TMQ_REV else '') for i in range(1, 19)]
    h += tmq('t0') + [f't0_se{i}' for i in range(1, 6)] + [f't0_kt{i}' for i in range(1, 9)] + ['t0_kt_diem']
    h += ['t1_thoi_gian'] + [f't1_kt{i}' for i in range(1, 9)] + ['t1_kt_diem']
    h += [f't1_se{i}' for i in range(1, 6)] + [f't1_re{i}' for i in range(1, 5)] + [f't1_yd{i}' for i in range(1, 4)]
    h += ['hd_cam', 'hd_dien_thoai', 'hd_hinh_1', 'hd_hinh_2'] + [f'hd_eisen_c{i}' for i in range(1, 9)]
    h += [f'hd_banhxe{i}' for i in range(1, 9)] + ['hd_muctieu_' + k for k in 'SMART']
    h += ['hd_muctieu_noidung', 'hd_camket', 'hd_danhgia_sao', 'hd_danhgia_noidung']
    for w in ('t2', 't3'):
        h += [f'{w}_thoi_gian', f'{w}_thamdu'] + [f'{w}_ad{i}' for i in range(1, 8)]
        if w == 't2':
            h += ['t2_camket_so_ngay']
        else:
            h += tmq('t3') + [f't3_se{i}' for i in range(1, 6)] + ['t3_thay_doi']
        h += [f'{w}_tro_ngai', f'{w}_y_kien']
    return h


# ----------------------------------------------------------------------------- định dạng số
DEC = ','


def fnum(x, k=2) -> str:
    if x is None or (isinstance(x, float) and (math.isnan(x) or math.isinf(x))):
        return '—'
    s = f'{x:.{k}f}'
    return s.replace('.', DEC) if DEC != '.' else s


def fint(x) -> str:
    return '—' if x is None or (isinstance(x, float) and math.isnan(x)) else str(int(x))


def fp(p) -> str:
    """p theo APA; với dấu phẩy thập phân giữ số 0 đứng đầu (thói quen tạp chí trong nước)."""
    if p is None or (isinstance(p, float) and math.isnan(p)):
        return '—'
    if p < .001:
        return '< ' + ('0,001' if DEC == ',' else '.001')
    s = f'{p:.3f}'
    if DEC == ',':
        return s.replace('.', ',')
    return s[1:] if s.startswith('0') else s


def fpct(x, k=1) -> str:
    return '—' if x is None or (isinstance(x, float) and math.isnan(x)) else fnum(100 * x, k) + '%'


# ----------------------------------------------------------------------------- thống kê
def cronbach_alpha(df: pd.DataFrame) -> float:
    d = df.dropna()
    k = d.shape[1]
    if len(d) < 3 or k < 2:
        return float('nan')
    var_items = d.var(ddof=1).sum()
    var_tot = d.sum(axis=1).var(ddof=1)
    return float(k / (k - 1) * (1 - var_items / var_tot)) if var_tot > 0 else float('nan')


def omega_total(df: pd.DataFrame) -> float:
    """McDonald ω từ mô hình một nhân tố (phân tích trục chính lặp trên ma trận tương quan)."""
    d = df.dropna()
    k = d.shape[1]
    if len(d) < max(10, k + 2):
        return float('nan')
    R = np.corrcoef(d.values.T)
    if not np.all(np.isfinite(R)):
        return float('nan')
    try:
        h2 = 1 - 1 / np.diag(np.linalg.inv(R))
    except np.linalg.LinAlgError:
        h2 = np.full(k, .5)
    for _ in range(200):
        Rr = R.copy()
        np.fill_diagonal(Rr, h2)
        w, v = np.linalg.eigh(Rr)
        lam = v[:, -1] * math.sqrt(max(w[-1], 1e-9))
        new = np.clip(lam ** 2, .001, .999)
        if np.max(np.abs(new - h2)) < 1e-6:
            break
        h2 = new
    lam = np.abs(lam) if lam.sum() < 0 else lam
    s = lam.sum() ** 2
    return float(s / (s + np.sum(1 - lam ** 2)))


def item_stats(df: pd.DataFrame) -> pd.DataFrame:
    d = df.dropna()
    rows = []
    for c in d.columns:
        rest = d.drop(columns=c)
        r = stats.pearsonr(d[c], rest.sum(axis=1))[0] if len(d) > 2 and d[c].std() > 0 else float('nan')
        rows.append({'Câu': c, 'M': d[c].mean(), 'SD': d[c].std(ddof=1), 'r biến–tổng hiệu chỉnh': r,
                     'α nếu bỏ câu': cronbach_alpha(rest)})
    return pd.DataFrame(rows)


def holm(ps: list[float]) -> list[float]:
    idx = [i for i, p in enumerate(ps) if p is not None and not math.isnan(p)]
    m = len(idx)
    out = [float('nan')] * len(ps)
    running = 0.0
    for rank, i in enumerate(sorted(idx, key=lambda j: ps[j])):
        running = max(running, min(1.0, (m - rank) * ps[i]))
        out[i] = running
    return out


def paired(a: pd.Series, b: pd.Series) -> dict:
    m = a.notna() & b.notna()
    x, y = a[m].astype(float), b[m].astype(float)
    n = int(m.sum())
    r = {'n': n, 'M1': x.mean(), 'SD1': x.std(ddof=1), 'M2': y.mean(), 'SD2': y.std(ddof=1)}
    if n < 3:
        return r
    d = y - x
    sd = d.std(ddof=1)
    r['dM'] = d.mean()
    se = sd / math.sqrt(n)
    tc = stats.t.ppf(.975, n - 1)
    r['CI_lo'], r['CI_hi'] = r['dM'] - tc * se, r['dM'] + tc * se
    if sd > 0:
        r['t'] = r['dM'] / se
        r['df'] = n - 1
        r['p_t'] = 2 * stats.t.sf(abs(r['t']), n - 1)
        r['dz'] = r['dM'] / sd
    nz = d[d != 0]
    if len(nz) >= 1:
        rk = stats.rankdata(nz.abs())
        wp, wm = rk[(nz > 0).values].sum(), rk[(nz < 0).values].sum()
        r['W'] = min(wp, wm)
        r['r_rb'] = (wp - wm) / (wp + wm)
        try:
            r['p_w'] = stats.wilcoxon(nz, zero_method='wilcox', alternative='two-sided').pvalue
        except ValueError:
            pass
    return r


def indep(x: pd.Series, y: pd.Series) -> dict:
    x, y = x.dropna().astype(float), y.dropna().astype(float)
    n1, n2 = len(x), len(y)
    r = {'n1': n1, 'M1': x.mean(), 'SD1': x.std(ddof=1), 'n2': n2, 'M2': y.mean(), 'SD2': y.std(ddof=1)}
    if n1 < 3 or n2 < 3:
        return r
    sp = math.sqrt(((n1 - 1) * x.var(ddof=1) + (n2 - 1) * y.var(ddof=1)) / (n1 + n2 - 2))
    if sp > 0:
        J = 1 - 3 / (4 * (n1 + n2) - 9)
        r['g'] = (x.mean() - y.mean()) / sp * J
    u = stats.mannwhitneyu(x, y, alternative='two-sided')
    r['U'], r['p'] = u.statistic, u.pvalue
    r['r_rb'] = 2 * u.statistic / (n1 * n2) - 1
    return r


def spearman(x: pd.Series, y: pd.Series) -> dict:
    m = x.notna() & y.notna()
    n = int(m.sum())
    if n < 5 or x[m].nunique() < 2 or y[m].nunique() < 2:
        return {'n': n}
    rho, p = stats.spearmanr(x[m], y[m])
    se = math.sqrt((1 + rho ** 2 / 2) / (n - 3))
    z = math.atanh(max(min(rho, .9999), -.9999))
    return {'n': n, 'rho': rho, 'p': p, 'lo': math.tanh(z - 1.96 * se), 'hi': math.tanh(z + 1.96 * se)}


def mcnemar(pre: pd.Series, post: pd.Series) -> dict:
    m = pre.notna() & post.notna()
    a, b = pre[m].astype(int), post[m].astype(int)
    up = int(((a == 0) & (b == 1)).sum())
    down = int(((a == 1) & (b == 0)).sum())
    p = stats.binomtest(min(up, down), up + down, .5).pvalue if up + down > 0 else float('nan')
    return {'n': int(m.sum()), 'p0': a.mean() if len(a) else float('nan'), 'p1': b.mean() if len(b) else float('nan'),
            'tang': up, 'giam': down, 'p': p}


def kruskal(groups: dict[str, pd.Series]) -> dict:
    gs = {k: v.dropna() for k, v in groups.items() if v.dropna().size >= 3}
    if len(gs) < 2:
        return {}
    H, p = stats.kruskal(*gs.values())
    n = sum(len(v) for v in gs.values())
    return {'H': H, 'df': len(gs) - 1, 'p': p, 'eps2': H * (n + 1) / (n ** 2 - 1), 'k': len(gs)}


# ----------------------------------------------------------------------------- đọc và chấm điểm
def load(path: str) -> pd.DataFrame:
    df = pd.read_csv(path, encoding='utf-8-sig', dtype=str, keep_default_na=False)
    df.columns = [c.strip() for c in df.columns]
    missing = [c for c in csv_header() if c not in df.columns]
    if missing:
        sys.exit(f'Tệp không đúng định dạng dữ liệu nghiên cứu — thiếu {len(missing)} cột, ví dụ: {missing[:5]}')
    text_cols = {'ma_an_danh', 't0_thoi_gian', 't1_thoi_gian', 't2_thoi_gian', 't3_thoi_gian', 'hd_hinh_1', 'hd_hinh_2',
                 'hd_muctieu_noidung', 'hd_danhgia_noidung', 't2_tro_ngai', 't3_tro_ngai', 't2_y_kien', 't3_y_kien'}
    for c in df.columns:
        if c not in text_cols:
            df[c] = pd.to_numeric(df[c].replace('', np.nan), errors='coerce')
    for c in ('co_t0', 'co_t1', 'co_t2', 'co_t3'):
        df[c] = df[c].fillna(0).astype(int)
    return df


def tmq_cols(w: str) -> list[str]:
    return [f'{w}_tmq{i:02d}_{TMQ_SUB[i - 1]}' + ('_R' if i in TMQ_REV else '') for i in range(1, 19)]


def score(df: pd.DataFrame) -> pd.DataFrame:
    s = df.copy()
    for w in ('t0', 't3'):
        raw = s[tmq_cols(w)].copy()
        for i in TMQ_REV:
            c = tmq_cols(w)[i - 1]
            raw[c] = 6 - raw[c]
        for i, c in enumerate(tmq_cols(w), start=1):
            s[f'{w}_q{i:02d}'] = raw[c]           # đã đảo chiều
        for sub in ('SRP', 'TA', 'LRP'):
            cols = [f'{w}_q{i:02d}' for i in range(1, 19) if TMQ_SUB[i - 1] == sub]
            s[f'{w}_{sub}'] = s[cols].mean(axis=1, skipna=False)
        s[f'{w}_TOT'] = s[[f'{w}_q{i:02d}' for i in range(1, 19)]].mean(axis=1, skipna=False)
        s[f'{w}_SE'] = s[[f'{w}_se{i}' for i in range(1, 6)]].mean(axis=1, skipna=False)
    s['t1_SE'] = s[[f't1_se{i}' for i in range(1, 6)]].mean(axis=1, skipna=False)
    for w in ('t0', 't1'):
        for i in range(1, 9):
            s[f'{w}_ok{i}'] = np.where(s[f'{w}_kt{i}'].isna(), np.nan, (s[f'{w}_kt{i}'] == KT_KEY[i - 1]).astype(float))
        s[f'{w}_KT'] = s[[f'{w}_ok{i}' for i in range(1, 9)]].sum(axis=1, min_count=8)
    s['t1_RE'] = s[[f't1_re{i}' for i in range(1, 5)]].mean(axis=1, skipna=False)
    s['t1_YD'] = s[[f't1_yd{i}' for i in range(1, 4)]].mean(axis=1, skipna=False)
    for w in ('t2', 't3'):
        ad = s[[f'{w}_ad{i}' for i in range(1, 8)]]
        s[f'{w}_AD'] = ad.mean(axis=1, skipna=False)
        s[f'{w}_AD_so_kt'] = (ad >= 1).sum(axis=1).where(ad.notna().all(axis=1))
        s[f'{w}_AD_deu'] = (ad >= 2).sum(axis=1).where(ad.notna().all(axis=1))
    s['t2_CK'] = s['t2_camket_so_ngay'].where(s['t2_camket_so_ngay'] <= 7)
    s['hd_SMART'] = s[['hd_muctieu_' + k for k in 'SMART']].sum(axis=1, min_count=5)
    ei = s[[f'hd_eisen_c{i}' for i in range(1, 9)]]
    s['hd_EISEN_dung'] = sum((ei[f'hd_eisen_c{i}'] == EISEN_KEY[f'c{i}']).astype(float) for i in range(1, 9)).where(ei.notna().all(axis=1))
    s['mot_kieu_t0'] = s[tmq_cols('t0')].nunique(axis=1).eq(1) & s[tmq_cols('t0')].notna().all(axis=1)
    s['tham_du'] = np.where(s['co_t0'].eq(1) | s['co_t1'].eq(1) | s['t2_thamdu'].eq(1) | s['t3_thamdu'].eq(1), 1, 0)
    return s


# ----------------------------------------------------------------------------- dữ liệu giả lập
def simulate(path: str, n0: int = 230, seed: int = 2026) -> None:
    rng = np.random.default_rng(seed)
    hdr = csv_header()
    used = set()

    def code():
        while True:
            c = f'ZZ{rng.integers(1, 32):02d}{rng.integers(0, 100):02d}'
            if c not in used:
                used.add(c)
                return c

    def lik(mu, sd=0.9):
        return int(np.clip(np.rint(rng.normal(mu, sd)), 1, 5))

    def kt(pcorr):
        out = []
        for j, p in enumerate(pcorr):
            if rng.random() < p:
                out.append(KT_KEY[j])
            else:
                out.append(int(rng.choice([o for o in range(4) if o != KT_KEY[j]])))
        return out

    base_t = dt.datetime(2026, 10, 18, 8, 5)
    rows = []
    p0 = np.array([.6, .4, .55, .65, .5, .45, .15, .2])
    p1 = np.array([.9, .85, .8, .92, .82, .88, .72, .6])
    texts = ['Trong 6 tuần, mỗi tối học 20 từ vựng tiếng Anh để đạt 6.0 bài thi thử 30/11',
             'Đạt từ 7 điểm Giải tích giữa kỳ bằng cách làm 10 bài tập mỗi tuần', 'Ngủ trước 23h ít nhất 5 đêm mỗi tuần',
             'Học tốt hơn', 'Mỗi Chủ nhật dành 20 phút lập kế hoạch tuần sau']
    for _ in range(n0):
        z = rng.normal()
        r = dict.fromkeys(hdr, '')
        r['ma_an_danh'] = code()
        r['so_thiet_bi'] = 1
        r['co_t0'] = 1
        dem = {'gioi': rng.choice(4, p=[.52, .45, .01, .02]), 'nganh': rng.integers(0, 10),
               'noio': rng.choice(5, p=[.45, .2, .28, .05, .02]), 'lamthem': rng.choice(4, p=[.6, .25, .12, .03]),
               'tinchi': rng.choice(5, p=[.1, .45, .3, .05, .1]), 'coso': rng.choice(3, p=[.7, .2, .1]),
               'dahoc': rng.choice(2, p=[.7, .3])}
        for k, v in dem.items():
            r['t0_' + k] = int(v)
        r['t0_thoi_gian'] = (base_t + dt.timedelta(minutes=int(rng.integers(0, 12)))).strftime('%Y-%m-%d %H:%M')
        mu = {'SRP': 2.7 + .55 * z, 'TA': 3.1 + .45 * z, 'LRP': 2.9 + .35 * z}
        t0q = []
        for i in range(1, 19):
            v = lik(mu[TMQ_SUB[i - 1]] + (.25 * z if i == 17 else 0), 1.0 if TMQ_SUB[i - 1] != 'SRP' else .8)
            t0q.append(6 - v if i in TMQ_REV else v)
        for c, v in zip(tmq_cols('t0'), t0q):
            r[c] = v
        se0 = [lik(3.0 + .45 * z) for _ in range(5)]
        for i, v in enumerate(se0, 1):
            r[f't0_se{i}'] = v
        bump = .12 * (dem['dahoc'] == 1)
        k0 = kt(np.clip(p0 + bump + .08 * z, .02, .98))
        for i, v in enumerate(k0, 1):
            r[f't0_kt{i}'] = v
        r['t0_kt_diem'] = sum(int(v == KT_KEY[j]) for j, v in enumerate(k0))
        hinh = rng.choice(list(SHAPES), p=[.38, .16, .06, .3, .1])
        r['hd_hinh_1'] = hinh
        r['hd_hinh_2'] = rng.choice([s for s in SHAPES if s != hinh]) if rng.random() < .7 else ''
        r['hd_cam'] = int(rng.choice(4, p=[.05, .35, .45, .15]))
        r['hd_dien_thoai'] = int(rng.choice(5, p=[.08, .27, .33, .2, .12]))
        for i in range(1, 9):
            r[f'hd_eisen_c{i}'] = EISEN_KEY[f'c{i}'] if rng.random() < .72 else int(rng.integers(1, 5))
        for i in range(1, 9):
            r[f'hd_banhxe{i}'] = int(np.clip(np.rint(rng.normal(6.2 + .6 * z, 1.8)), 1, 10))
        for k, p in zip('SMART', [.8, .55, .75, .85, .6]):
            r['hd_muctieu_' + k] = int(rng.random() < p)
        r['hd_muctieu_noidung'] = str(rng.choice(texts))
        r['hd_camket'] = int(rng.choice(8, p=[.2, .18, .1, .17, .12, .1, .08, .05]))
        r['hd_danhgia_sao'] = int(rng.choice([3, 4, 5], p=[.1, .4, .5]))
        r['hd_danhgia_noidung'] = str(rng.choice(['Hoạt động trên điện thoại rất vui', 'Muốn thêm thời gian hỏi đáp', '']))
        if rng.random() < .86:
            r['co_t1'] = 1
            r['t1_thoi_gian'] = (base_t + dt.timedelta(minutes=110 + int(rng.integers(0, 10)))).strftime('%Y-%m-%d %H:%M')
            k1 = kt(np.clip(p1 + .05 * z, .02, .98))
            for i, v in enumerate(k1, 1):
                r[f't1_kt{i}'] = v
            r['t1_kt_diem'] = sum(int(v == KT_KEY[j]) for j, v in enumerate(k1))
            for i, v in enumerate(se0, 1):
                r[f't1_se{i}'] = int(np.clip(v + rng.choice([0, 0, 1, 1, 1, 2]) - (rng.random() < .08), 1, 5))
            for i in range(1, 5):
                r[f't1_re{i}'] = lik(4.2, .7)
            yd = [lik(4.0 + .2 * z, .8) for _ in range(3)]
            for i, v in enumerate(yd, 1):
                r[f't1_yd{i}'] = v
            yd_m = np.mean(yd)
        else:
            r['co_t1'] = 0
            yd_m = 3.8
        apply_lvl = .9 + .35 * (yd_m - 4) + .3 * z
        if rng.random() < .46:
            r['co_t2'] = 1
            r['t2_thoi_gian'] = f'2026-11-0{int(rng.integers(1, 10))} 2{int(rng.integers(0, 3))}:15'
            r['t2_thamdu'] = 1
            for i in range(1, 8):
                r[f't2_ad{i}'] = int(np.clip(np.rint(rng.normal(apply_lvl + [.2, 0, .3, -.1, .4, .3, -.2][i - 1], .9)), 0, 3))
            r['t2_camket_so_ngay'] = int(rng.choice(9, p=[.05, .06, .09, .12, .14, .15, .12, .17, .1]))
            r['t2_tro_ngai'] = ';'.join(str(x) for x in sorted(rng.choice(8, size=int(rng.integers(0, 3)), replace=False)))
            r['t2_y_kien'] = str(rng.choice(['', '', 'Pomodoro hợp với mình', 'Khó duy trì khi lịch thi dày']))
        if rng.random() < .39:
            r['co_t3'] = 1
            r['t3_thoi_gian'] = f'2026-12-0{int(rng.integers(1, 10))} 20:40'
            r['t3_thamdu'] = 1
            for i in range(1, 8):
                r[f't3_ad{i}'] = int(np.clip(np.rint(rng.normal(apply_lvl - .15 + [.2, 0, .3, -.1, .4, .3, -.2][i - 1], .9)), 0, 3))
            gain = {'SRP': .30, 'TA': .10, 'LRP': .12}
            for i, c in enumerate(tmq_cols('t3'), start=1):
                v = lik(mu[TMQ_SUB[i - 1]] + gain[TMQ_SUB[i - 1]], 1.0 if TMQ_SUB[i - 1] != 'SRP' else .8)
                r[c] = 6 - v if i in TMQ_REV else v
            for i in range(1, 6):
                r[f't3_se{i}'] = lik(3.35 + .45 * z)
            r['t3_thay_doi'] = int(rng.choice([2, 3, 4, 5], p=[.05, .3, .5, .15]))
            r['t3_tro_ngai'] = ';'.join(str(x) for x in sorted(rng.choice(8, size=int(rng.integers(0, 4)), replace=False)))
        rows.append(r)
    # người không tham dự, chỉ trả lời T3 (nhóm so sánh)
    for _ in range(int(n0 * .2)):
        z = rng.normal() - .15
        r = dict.fromkeys(hdr, '')
        r.update({'ma_an_danh': code(), 'so_thiet_bi': 0, 'co_t0': 0, 'co_t1': 0, 'co_t2': 0, 'co_t3': 1})
        r['t3_thoi_gian'] = f'2026-12-0{int(rng.integers(1, 10))} 21:05'
        r['t3_thamdu'] = 0
        for i in range(1, 8):
            r[f't3_ad{i}'] = int(np.clip(np.rint(rng.normal(.55 + .3 * z, .8)), 0, 3))
        mu = {'SRP': 2.75 + .55 * z, 'TA': 3.1 + .45 * z, 'LRP': 2.9 + .35 * z}
        for i, c in enumerate(tmq_cols('t3'), start=1):
            v = lik(mu[TMQ_SUB[i - 1]], 1.0 if TMQ_SUB[i - 1] != 'SRP' else .8)
            r[c] = 6 - v if i in TMQ_REV else v
        for i in range(1, 6):
            r[f't3_se{i}'] = lik(3.05 + .45 * z)
        r['t3_thay_doi'] = int(rng.choice([2, 3, 4, 5], p=[.08, .5, .35, .07]))
        r['t3_tro_ngai'] = ';'.join(str(x) for x in sorted(rng.choice(8, size=int(rng.integers(0, 3)), replace=False)))
        rows.append(r)
    out = pd.DataFrame(rows, columns=hdr)
    out.to_csv(path, index=False, encoding='utf-8-sig', lineterminator='\r\n')
    print(f'Đã tạo dữ liệu GIẢ LẬP: {path} ({len(out)} dòng). Không dùng tệp này cho bài báo.')


# ----------------------------------------------------------------------------- phân tích
class Report:
    """Gom bảng (cho Excel + Word), hình và câu kết quả."""

    def __init__(self):
        self.blocks: list[tuple] = []      # ('h', text) | ('p', text) | ('t', title, df, note) | ('img', path, caption)
        self.sheets: dict[str, pd.DataFrame] = {}
        self.ntab = 0
        self.nfig = 0

    def h(self, text):
        self.blocks.append(('h', text))

    def p(self, text):
        self.blocks.append(('p', text))

    def table(self, title, df: pd.DataFrame, sheet: str, note: str = '', raw: pd.DataFrame | None = None):
        self.ntab += 1
        self.blocks.append(('t', f'Bảng {self.ntab}. {title}', df, note))
        self.sheets[f'B{self.ntab:02d}_{sheet}'[:31]] = raw if raw is not None else df

    def fig(self, path, caption):
        self.nfig += 1
        self.blocks.append(('img', path, f'Hình {self.nfig}. {caption}'))


def apa_t(r: dict) -> str:
    if 't' not in r:
        return 'không đủ dữ liệu để kiểm định'
    return (f't({fint(r["df"])}) = {fnum(r["t"])}, p {"= " if r["p_t"] >= .001 else ""}{fp(r["p_t"])}, '
            f'd_z = {fnum(r["dz"])}, 95% KTC của chênh lệch [{fnum(r["CI_lo"])}; {fnum(r["CI_hi"])}]')


def analyse(df: pd.DataFrame, s: pd.DataFrame, outdir: str, synthetic: bool, src: str, drop_straight: bool) -> Report:
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.family': 'DejaVu Sans', 'font.size': 10, 'axes.spines.top': False,
                         'axes.spines.right': False, 'figure.dpi': 150})
    figdir = os.path.join(outdir, 'bieu_do')
    os.makedirs(figdir, exist_ok=True)

    def savefig(fig, name):
        if synthetic:
            fig.text(.5, .5, 'DỮ LIỆU GIẢ LẬP', ha='center', va='center', fontsize=34, color='#D2492A', alpha=.13, rotation=20)
        p = os.path.join(figdir, name)
        fig.savefig(p, bbox_inches='tight')
        plt.close(fig)
        return p

    R = Report()
    t0 = s[s.co_t0 == 1]
    if drop_straight:
        t0_tmq = t0[~t0.mot_kieu_t0]
    else:
        t0_tmq = t0

    # ---------------- 1. Mẫu
    R.h('1. Mẫu và dòng tham gia')
    att = s[s.tham_du == 1]
    flow = [
        ('Mã ẩn danh trong tệp', len(s)),
        ('Có khảo sát đầu buổi (T0)', int(s.co_t0.sum())),
        ('Có khảo sát cuối buổi (T1)', int(s.co_t1.sum())),
        ('Có cả T0 và T1', int((s.co_t0 & s.co_t1).sum())),
        ('Có khảo sát sau 2 tuần (T2), người tham dự', int(((s.co_t2 == 1) & (s.t2_thamdu == 1)).sum())),
        ('Có khảo sát sau 6 tuần (T3), người tham dự', int(((s.co_t3 == 1) & (s.t3_thamdu == 1)).sum())),
        ('Có cả T0 và T3 (người tham dự)', int(((s.co_t0 == 1) & (s.co_t3 == 1) & (s.t3_thamdu == 1)).sum())),
        ('T3 — người KHÔNG tham dự (nhóm so sánh)', int(((s.co_t3 == 1) & (s.t3_thamdu == 0)).sum())),
        ('Mã xuất hiện trên > 1 thiết bị (khả năng trùng mã)', int((s.so_thiet_bi > 1).sum())),
        ('Trả lời TMQ một kiểu ở T0 (18 câu cùng mức)', int(t0.mot_kieu_t0.sum())),
    ]
    n_t0 = max(int(s.co_t0.sum()), 1)
    fdf = pd.DataFrame([{'Nhóm': a, 'n': b, '% so với T0': fpct(b / n_t0) if i and i < 8 else ''} for i, (a, b) in enumerate(flow)])
    R.table('Dòng tham gia qua các đợt khảo sát', fdf, 'dong_tham_gia',
            note='Tỷ lệ tính trên số người có T0. ' + ('Các phân tích TMQ ở T0 đã loại người trả lời một kiểu.' if drop_straight else
                                                         'Người trả lời một kiểu được giữ lại; chạy lại với --loai-mot-kieu để phân tích độ nhạy.'))
    rows = []
    for k, (lab, opts) in DEMO.items():
        vc = t0['t0_' + k].value_counts()
        for i, o in enumerate(opts):
            n = int(vc.get(i, 0))
            rows.append({'Biến': lab, 'Nhóm': o, 'n': n, '%': fpct(n / len(t0)) if len(t0) else '—'})
    R.table(f'Đặc điểm mẫu tại T0 (n = {len(t0)})', pd.DataFrame(rows), 'dac_diem_mau')

    # ---------------- 2. Độ tin cậy
    R.h('2. Độ tin cậy của thang đo')
    rel = []
    item_tabs = {}
    specs = []
    for w, frame in (('T0', t0_tmq), ('T3', s[s.co_t3 == 1])):
        pre = w.lower()
        for sub in ('SRP', 'TA', 'LRP'):
            specs.append((f'TMQ — {TMQ_NAME[sub]} ({w})', frame[[f'{pre}_q{i:02d}' for i in range(1, 19) if TMQ_SUB[i - 1] == sub]]))
        specs.append((f'TMQ — toàn thang ({w})', frame[[f'{pre}_q{i:02d}' for i in range(1, 19)]]))
        specs.append((f'Tự tin QLTG ({w})', frame[[f'{pre}_se{i}' for i in range(1, 6)]]))
    specs += [('Tự tin QLTG (T1)', s[[f't1_se{i}' for i in range(1, 6)]]),
              ('Phản hồi về workshop (T1)', s[[f't1_re{i}' for i in range(1, 5)]]),
              ('Ý định vận dụng (T1)', s[[f't1_yd{i}' for i in range(1, 4)]]),
              ('Mức vận dụng 7 kỹ thuật (T2)', s[[f't2_ad{i}' for i in range(1, 8)]]),
              ('Mức vận dụng 7 kỹ thuật (T3)', s[[f't3_ad{i}' for i in range(1, 8)]])]
    for name, frame in specs:
        d = frame.dropna()
        rel.append({'Thang đo': name, 'Số câu': frame.shape[1], 'n': len(d), 'α': fnum(cronbach_alpha(frame)),
                    'ω': fnum(omega_total(frame))})
    item_tabs['TMQ_T0'] = item_stats(t0_tmq[[f't0_q{i:02d}' for i in range(1, 19)]])
    R.table('Hệ số tin cậy Cronbach α và McDonald ω', pd.DataFrame(rel), 'do_tin_cay',
            note='ω tính từ mô hình một nhân tố cho từng thang. Thang có α < 0,60 cần thận trọng khi diễn giải; '
                 'cấu trúc 3 nhân tố của TMQ kiểm định bằng CFA (tệp cfa_tmq.R).')
    it = item_tabs['TMQ_T0'].copy()
    it.insert(1, 'Thang con', [TMQ_SUB[i] + (' (đảo)' if i + 1 in TMQ_REV else '') for i in range(18)])
    it_fmt = it.copy()
    for c in ('M', 'SD', 'r biến–tổng hiệu chỉnh', 'α nếu bỏ câu'):
        it_fmt[c] = it[c].map(fnum)
    it_fmt['Câu'] = [f'Câu {i}' for i in range(1, 19)]
    weak = [f'câu {i + 1}' for i, v in enumerate(it['r biến–tổng hiệu chỉnh']) if not math.isnan(v) and v < .20]
    R.table('Thống kê từng câu TMQ tại T0 (đã đảo chiều)', it_fmt, 'tmq_tung_cau',
            note=('Câu có tương quan biến–tổng hiệu chỉnh < 0,20: ' + ', '.join(weak) + '. Cân nhắc báo cáo phân tích có và không có các câu này.')
            if weak else 'Mọi câu có tương quan biến–tổng hiệu chỉnh ≥ 0,20.', raw=it)

    # ---------------- 3. Thực trạng (RQ1)
    R.h('3. Thực trạng kỹ năng quản lý thời gian tại T0 (CH1)')
    desc = []
    for k in ('SRP', 'TA', 'LRP', 'TOT'):
        v = t0_tmq[f't0_{k}'].dropna()
        desc.append({'Thang': TMQ_NAME[k], 'n': len(v), 'M': fnum(v.mean()), 'SD': fnum(v.std(ddof=1)),
                     'Trung vị': fnum(v.median()), '% điểm TB ≥ 3,5': fpct((v >= 3.5).mean()) if len(v) else '—'})
    for k, lab in (('t0_SE', 'Tự tin QLTG'), ('t0_KT', 'Điểm kiến thức (0–8)')):
        v = t0[k].dropna()
        desc.append({'Thang': lab, 'n': len(v), 'M': fnum(v.mean()), 'SD': fnum(v.std(ddof=1)), 'Trung vị': fnum(v.median()),
                     '% điểm TB ≥ 3,5': fpct((v >= 3.5).mean()) if k == 't0_SE' and len(v) else ''})
    R.table('Thống kê mô tả các thang đo tại T0', pd.DataFrame(desc), 'mo_ta_t0', note='TMQ và tự tin: thang 1–5, điểm cao = tốt hơn.')

    grp_rows = []
    for k in ('gioi', 'lamthem', 'noio', 'dahoc', 'tinchi'):
        lab, opts = DEMO[k]
        for sc in ('TOT', 'SRP'):
            groups = {opts[i]: t0_tmq.loc[t0_tmq['t0_' + k] == i, f't0_{sc}'] for i in range(len(opts))
                      if opts[i] not in ('Không muốn trả lời', 'Không rõ', 'Khác')}
            kr = kruskal(groups)
            means = '; '.join(f'{g}: {fnum(v.mean())} (n={v.notna().sum()})' for g, v in groups.items() if v.notna().sum() >= 3)
            grp_rows.append({'Biến nhóm': lab, 'Thang': TMQ_NAME[sc], 'Trung bình theo nhóm': means or '—',
                             'H': fnum(kr.get('H')), 'df': fint(kr.get('df')), 'p': fp(kr.get('p')), 'ε²': fnum(kr.get('eps2'), 3)})
    R.table('So sánh TMQ tại T0 theo đặc điểm nền (Kruskal–Wallis)', pd.DataFrame(grp_rows), 'tmq_theo_nhom',
            note='Phân tích khám phá; chỉ tính nhóm có ≥ 3 người. ε² ≈ 0,01 nhỏ; 0,08 trung bình; 0,26 lớn.')
    cor = []
    for a, la in (('t0_SRP', TMQ_NAME['SRP']), ('t0_TA', TMQ_NAME['TA']), ('t0_LRP', TMQ_NAME['LRP']), ('t0_TOT', TMQ_NAME['TOT'])):
        for b, lb in (('t0_SE', 'Tự tin QLTG (T0)'), ('t0_KT', 'Kiến thức (T0)'), ('hd_dien_thoai', 'Thời gian dùng điện thoại'),
                      ('hd_cam', 'Cảm nhận áp lực thời gian')):
            r = spearman(t0_tmq[a], t0_tmq[b])
            cor.append({'Biến 1': la, 'Biến 2': lb, 'n': r['n'], 'ρ': fnum(r.get('rho')), '95% KTC': (f'[{fnum(r["lo"])}; {fnum(r["hi"])}]' if 'lo' in r else '—'),
                        'p': fp(r.get('p'))})
    R.table('Tương quan Spearman giữa TMQ và các biến khác tại T0', pd.DataFrame(cor), 'tuong_quan_t0',
            note='Thời gian dùng điện thoại và cảm nhận áp lực lấy từ hoạt động Khởi động (thang thứ bậc).')
    if len(t0_tmq):
        fig, ax = plt.subplots(figsize=(6.4, 3.2))
        data = [t0_tmq[f't0_{k}'].dropna() for k in ('SRP', 'TA', 'LRP')]
        if all(len(d) for d in data):
            try:
                bp = ax.boxplot(data, orientation='horizontal', patch_artist=True, widths=.55)
            except TypeError:   # matplotlib < 3.10
                bp = ax.boxplot(data, vert=False, patch_artist=True, widths=.55)
            for patch, c in zip(bp['boxes'], (MAU['luc'], MAU['lam'], MAU['vang'])):
                patch.set_facecolor(c)
                patch.set_alpha(.35)
            ax.set_yticks([1, 2, 3], [TMQ_NAME[k] for k in ('SRP', 'TA', 'LRP')])
            ax.set_xlim(1, 5)
            ax.set_xlabel('Điểm trung bình (1–5)')
            R.fig(savefig(fig, '01_tmq_t0.png'), f'Phân bố điểm ba thang con TMQ tại T0 (n = {len(t0_tmq)})')
        else:
            plt.close(fig)

    # ---------------- 4. Kết quả học tập (H1)
    R.h('4. Kết quả học tập ngay sau workshop — T0 so với T1 (GT1)')
    pr = []
    tests = [('Điểm kiến thức (0–8)', 't0_KT', 't1_KT'), ('Tự tin QLTG (1–5)', 't0_SE', 't1_SE')]
    tests += [(f'  Tự tin: {SE_NAME[i - 1]}', f't0_se{i}', f't1_se{i}') for i in range(1, 6)]
    res = [paired(s[a], s[b]) for _, a, b in tests]
    adj = holm([res[0].get('p_t', float('nan')), res[1].get('p_t', float('nan'))]) + [float('nan')] * 5
    for (lab, _, _), r, pa in zip(tests, res, adj):
        pr.append({'Biến': lab, 'n': r['n'], 'T0 M (SD)': f'{fnum(r["M1"])} ({fnum(r["SD1"])})',
                   'T1 M (SD)': f'{fnum(r["M2"])} ({fnum(r["SD2"])})', 'Chênh lệch [95% KTC]':
                       f'{fnum(r.get("dM"))} [{fnum(r.get("CI_lo"))}; {fnum(r.get("CI_hi"))}]',
                   't': fnum(r.get('t')), 'p': fp(r.get('p_t')), 'p (Holm)': fp(pa), 'd_z': fnum(r.get('dz')),
                   'Wilcoxon p': fp(r.get('p_w')), 'r_rb': fnum(r.get('r_rb'))})
    R.table('So sánh trước–sau (kiểm định t ghép cặp và Wilcoxon)', pd.DataFrame(pr), 'truoc_sau_t0_t1',
            note='Hai biến chính (kiến thức, tự tin tổng) hiệu chỉnh Holm; các câu tự tin là phân tích bổ sung. '
                 'd_z = chênh lệch trung bình / độ lệch chuẩn của chênh lệch; r_rb = tương quan hạng–nhị phân (Wilcoxon).')
    kt, se = res[0], res[1]
    if 't' in kt:
        R.p(f'Câu kết quả (gợi ý): Điểm kiến thức tăng từ M = {fnum(kt["M1"])} (SD = {fnum(kt["SD1"])}) ở đầu buổi lên '
            f'M = {fnum(kt["M2"])} (SD = {fnum(kt["SD2"])}) ở cuối buổi, {apa_t(kt)}.')
    if 't' in se:
        R.p(f'Câu kết quả (gợi ý): Mức tự tin về quản lý thời gian thay đổi từ M = {fnum(se["M1"])} (SD = {fnum(se["SD1"])}) '
            f'lên M = {fnum(se["M2"])} (SD = {fnum(se["SD2"])}), {apa_t(se)}.')
    mc = []
    mres = [mcnemar(s[f't0_ok{i}'], s[f't1_ok{i}']) for i in range(1, 9)]
    madj = holm([m['p'] for m in mres])
    for i, (m, pa) in enumerate(zip(mres, madj)):
        mc.append({'Câu': KT_NAME[i], 'n': m['n'], '% đúng T0': fpct(m['p0']), '% đúng T1': fpct(m['p1']),
                   'Sai→Đúng': m['tang'], 'Đúng→Sai': m['giam'], 'p (McNemar chính xác)': fp(m['p']), 'p (Holm)': fp(pa)})
    R.table('Tỷ lệ trả lời đúng từng câu kiến thức, T0 so với T1', pd.DataFrame(mc), 'mcnemar_kien_thuc',
            note='Chỉ tính người có cả T0 và T1. Kiểm định McNemar chính xác (nhị thức), hiệu chỉnh Holm cho 8 câu.')
    both = s[(s.co_t0 == 1) & (s.co_t1 == 1)]
    if len(both) >= 3:
        fig, axs = plt.subplots(1, 2, figsize=(11, 3.8), gridspec_kw={'width_ratios': [1, 1.4]}, constrained_layout=True)
        bins = np.arange(0, 9)
        c0 = both['t0_KT'].value_counts().reindex(bins, fill_value=0)
        c1 = both['t1_KT'].value_counts().reindex(bins, fill_value=0)
        axs[0].bar(bins - .2, c0 / len(both) * 100, width=.4, color=MAU['xam'], label='Đầu buổi (T0)')
        axs[0].bar(bins + .2, c1 / len(both) * 100, width=.4, color=MAU['luc'], label='Cuối buổi (T1)')
        axs[0].set_xticks(bins)
        axs[0].set_xlabel('Số câu đúng (0–8)')
        axs[0].set_ylabel('% người tham gia')
        axs[0].legend(frameon=False)
        y = np.arange(8)
        axs[1].barh(y - .2, [m['p0'] * 100 for m in mres], height=.4, color=MAU['xam'])
        axs[1].barh(y + .2, [m['p1'] * 100 for m in mres], height=.4, color=MAU['luc'])
        axs[1].set_yticks(y, [k.split('. ')[1] for k in KT_NAME])
        axs[1].invert_yaxis()
        axs[1].set_xlim(0, 100)
        axs[1].set_xlabel('% trả lời đúng')
        R.fig(savefig(fig, '02_kien_thuc_t0_t1.png'), f'Kiến thức trước và sau workshop (n = {len(both)} người có cả T0 và T1)')

    # ---------------- 5. Phản hồi
    R.h('5. Phản hồi về workshop và ý định vận dụng (Kirkpatrick mức 1)')
    fb = []
    for pre, names in (('re', RE_NAME), ('yd', YD_NAME)):
        for i, nm in enumerate(names, 1):
            v = s[f't1_{pre}{i}'].dropna()
            fb.append({'Nội dung': ('Phản hồi — ' if pre == 're' else 'Ý định — ') + nm, 'n': len(v), 'M': fnum(v.mean()),
                       'SD': fnum(v.std(ddof=1)), '% đồng ý (4–5)': fpct((v >= 4).mean()) if len(v) else '—'})
    sao = s['hd_danhgia_sao'].dropna()
    fb.append({'Nội dung': 'Số sao đánh giá chung (1–5, hoạt động Góp ý)', 'n': len(sao), 'M': fnum(sao.mean()),
               'SD': fnum(sao.std(ddof=1)), '% đồng ý (4–5)': fpct((sao >= 4).mean()) if len(sao) else '—'})
    R.table('Phản hồi và ý định vận dụng ở cuối buổi (T1)', pd.DataFrame(fb), 'phan_hoi_y_dinh', note='Thang 1–5.')

    # ---------------- 6. Vận dụng (H2)
    R.h('6. Mức độ vận dụng sau 2 và 6 tuần (GT2, Kirkpatrick mức 3)')
    att2 = s[(s.co_t2 == 1) & (s.t2_thamdu == 1)]
    att3 = s[(s.co_t3 == 1) & (s.t3_thamdu == 1)]
    ad = []
    for i, nm in enumerate(AD_NAME, 1):
        row = {'Kỹ thuật': nm}
        for w, fr in (('T2', att2), ('T3', att3)):
            v = fr[f'{w.lower()}_ad{i}'].dropna()
            row[f'{w}: % đã dùng'] = fpct((v >= 1).mean()) if len(v) else '—'
            row[f'{w}: % dùng đều'] = fpct((v >= 2).mean()) if len(v) else '—'
            row[f'{w}: M'] = fnum(v.mean())
        ad.append(row)
    R.table(f'Tỷ lệ vận dụng từng kỹ thuật ở người tham dự (T2: n = {len(att2)}; T3: n = {len(att3)})', pd.DataFrame(ad),
            'van_dung', note='Thang 0 = chưa lần nào, 1 = 1–2 lần, 2 = vài lần mỗi tuần, 3 = gần như mỗi ngày. '
                             '“Đã dùng” = ≥ 1; “dùng đều” = ≥ 2 (vài lần mỗi tuần trở lên).')
    vv = []
    for lab, a, b in (('Ý định (T1) → mức vận dụng (T2)', 't1_YD', 't2_AD'), ('Ý định (T1) → mức vận dụng (T3)', 't1_YD', 't3_AD'),
                      ('Kiến thức (T1) → mức vận dụng (T2)', 't1_KT', 't2_AD'), ('Tự tin (T1) → mức vận dụng (T2)', 't1_SE', 't2_AD'),
                      ('Số tiêu chí SMART tự kiểm → mức vận dụng (T2)', 'hd_SMART', 't2_AD'),
                      ('Số ngày giữ cam kết (T2) → mức vận dụng (T3)', 't2_CK', 't3_AD')):
        r = spearman(s[a], s[b])
        vv.append({'Quan hệ': lab, 'n': r['n'], 'ρ': fnum(r.get('rho')),
                   '95% KTC': f'[{fnum(r["lo"])}; {fnum(r["hi"])}]' if 'lo' in r else '—', 'p': fp(r.get('p'))})
    R.table('Tương quan giữa ý định, kết quả học tập và mức vận dụng (Spearman)', pd.DataFrame(vv), 'y_dinh_van_dung',
            note='Kiểm định giả thuyết GT2 (ý định dự báo hành vi, theo Thuyết hành vi có kế hoạch).')
    decay = paired(s.loc[s.t2_thamdu.eq(1), 't2_AD'], s.loc[s.t2_thamdu.eq(1), 't3_AD'])
    R.p(f'Mức vận dụng trung bình T2 so với T3 (cùng người, n = {decay["n"]}): T2 M = {fnum(decay["M1"])}, '
        f'T3 M = {fnum(decay["M2"])}; {apa_t(decay)}; Wilcoxon p = {fp(decay.get("p_w"))}.')
    ck = att2['t2_camket_so_ngay'].dropna()
    ck_rows = [{'Số ngày giữ cam kết 7 ngày': (f'{i} ngày' if i < 7 else 'Cả 7 ngày') if i < 8 else 'Không chọn cam kết',
                'n': int((ck == i).sum()), '%': fpct((ck == i).mean()) if len(ck) else '—'} for i in range(9)]
    R.table('Số ngày thực hiện cam kết 7 ngày (T2)', pd.DataFrame(ck_rows), 'cam_ket_7_ngay')
    tr_rows = []
    for w, fr in (('T2', att2), ('T3', att3)):
        col = fr[f'{w.lower()}_tro_ngai'].fillna('')
        n = len(fr)
        for i, nm in enumerate(TR_NAME):
            k = sum(1 for x in col if str(i) in str(x).split(';'))
            tr_rows.append({'Đợt': w, 'Trở ngại': nm, 'n': k, '%': fpct(k / n) if n else '—'})
    trdf = pd.DataFrame(tr_rows)
    R.table('Trở ngại khi vận dụng (chọn nhiều)', trdf, 'tro_ngai')
    if len(att2) + len(att3) >= 3:
        fig, axs = plt.subplots(1, 2, figsize=(11, 3.8), constrained_layout=True)
        y = np.arange(7)
        for w, fr, off, c in (('T2', att2, -.2, MAU['lam']), ('T3', att3, .2, MAU['vang'])):
            vals = [((fr[f'{w.lower()}_ad{i}'] >= 2).mean() * 100) if fr[f'{w.lower()}_ad{i}'].notna().any() else 0 for i in range(1, 8)]
            axs[0].barh(y + off, vals, height=.4, color=c, label=f'{w} (n = {len(fr)})')
        axs[0].set_yticks(y, AD_NAME)
        axs[0].invert_yaxis()
        axs[0].set_xlim(0, 100)
        axs[0].set_xlabel('% dùng vài lần mỗi tuần trở lên')
        fig.legend(frameon=False, loc='outside upper center', ncol=2)
        for w, off, c in (('T2', -.2, MAU['lam']), ('T3', .2, MAU['vang'])):
            sub = trdf[trdf['Đợt'] == w]
            nfr = len(att2 if w == 'T2' else att3) or 1
            axs[1].barh(np.arange(8) + off, sub['n'] / nfr * 100, height=.4, color=c)
        axs[1].set_yticks(np.arange(8), TR_NAME)
        axs[1].invert_yaxis()
        axs[1].set_xlabel('% người chọn')
        R.fig(savefig(fig, '03_van_dung_tro_ngai.png'), 'Mức vận dụng thường xuyên từng kỹ thuật và các trở ngại, sau 2 tuần (T2) và 6 tuần (T3)')

    # ---------------- 7. Thay đổi TMQ (H3)
    R.h('7. Thay đổi thói quen quản lý thời gian sau 6 tuần — T0 so với T3 (GT3)')
    base = s[(s.co_t0 == 1) & (s.co_t3 == 1) & (s.t3_thamdu == 1)]
    if drop_straight:
        base = base[~base.mot_kieu_t0]
    rows3, res3 = [], []
    for k in ('SRP', 'TA', 'LRP', 'TOT', 'SE'):
        res3.append(paired(base[f't0_{k}'], base[f't3_{k}']))
    adj3 = holm([r.get('p_t', float('nan')) for r in res3[:3]]) + [float('nan'), float('nan')]
    for k, r, pa in zip(('SRP', 'TA', 'LRP', 'TOT', 'SE'), res3, adj3):
        rows3.append({'Thang': TMQ_NAME.get(k, 'Tự tin QLTG'), 'n': r['n'], 'T0 M (SD)': f'{fnum(r["M1"])} ({fnum(r["SD1"])})',
                      'T3 M (SD)': f'{fnum(r["M2"])} ({fnum(r["SD2"])})',
                      'Chênh lệch [95% KTC]': f'{fnum(r.get("dM"))} [{fnum(r.get("CI_lo"))}; {fnum(r.get("CI_hi"))}]',
                      't': fnum(r.get('t')), 'p': fp(r.get('p_t')), 'p (Holm)': fp(pa), 'd_z': fnum(r.get('dz')),
                      'Wilcoxon p': fp(r.get('p_w'))})
    R.table('Thay đổi TMQ và tự tin từ T0 đến T3 ở người tham dự', pd.DataFrame(rows3), 'tmq_t0_t3',
            note='Ba thang con hiệu chỉnh Holm. Thiết kế một nhóm không loại trừ được tác động của thời gian (trưởng thành, '
                 'quen dần môi trường đại học); xem so sánh với nhóm không tham dự ở bảng sau.')
    if 't' in res3[0]:
        R.p(f'Câu kết quả (gợi ý): Sau 6 tuần, điểm lập kế hoạch ngắn hạn thay đổi từ M = {fnum(res3[0]["M1"])} lên '
            f'M = {fnum(res3[0]["M2"])}, {apa_t(res3[0])}.')
    td = att3['t3_thay_doi'].dropna()
    R.p('Tự đánh giá mức thay đổi trong quản lý thời gian (T3, người tham dự): ' +
        '; '.join(f'{nm}: {fpct((td == i).mean()) if len(td) else "—"}' for i, nm in enumerate(THAYDOI, 1)) + f' (n = {len(td)}).')
    non = s[(s.co_t3 == 1) & (s.t3_thamdu == 0)]
    cmp_rows = []
    for k, lab in (('t3_SRP', TMQ_NAME['SRP']), ('t3_TA', TMQ_NAME['TA']), ('t3_LRP', TMQ_NAME['LRP']), ('t3_TOT', TMQ_NAME['TOT']),
                   ('t3_SE', 'Tự tin QLTG'), ('t3_AD', 'Mức vận dụng 7 kỹ thuật (0–3)'), ('t3_thay_doi', 'Tự đánh giá thay đổi (1–5)')):
        r = indep(att3[k], non[k])
        cmp_rows.append({'Biến (T3)': lab, 'Tham dự n': r['n1'], 'Tham dự M (SD)': f'{fnum(r["M1"])} ({fnum(r["SD1"])})',
                         'Không tham dự n': r['n2'], 'Không tham dự M (SD)': f'{fnum(r["M2"])} ({fnum(r["SD2"])})',
                         'Hedges g': fnum(r.get('g')), 'U': fnum(r.get('U'), 0), 'p': fp(r.get('p')), 'r_rb': fnum(r.get('r_rb'))})
    R.table('So sánh người tham dự và không tham dự tại T3 (Mann–Whitney)', pd.DataFrame(cmp_rows), 'tham_du_vs_khong',
            note='Hai nhóm không được phân ngẫu nhiên (tự chọn tham dự), nên chênh lệch chỉ mang tính tham khảo, không phải bằng chứng nhân quả.')
    if len(base) >= 3:
        fig, ax = plt.subplots(figsize=(6.4, 3.3))
        ks = ('SRP', 'TA', 'LRP')
        x = np.arange(3)
        for off, w, c, lab in ((-.2, 't0', MAU['xam'], 'Đầu buổi (T0)'), (.2, 't3', MAU['luc'], 'Sau 6 tuần (T3)')):
            m = [base[f'{w}_{k}'].mean() for k in ks]
            e = [1.96 * base[f'{w}_{k}'].std(ddof=1) / math.sqrt(len(base)) for k in ks]
            ax.bar(x + off, m, width=.4, color=c, yerr=e, capsize=3, label=lab)
        ax.set_xticks(x, [TMQ_NAME[k] for k in ks])
        ax.set_ylim(1, 5)
        ax.set_ylabel('Điểm trung bình (1–5)')
        ax.legend(frameon=False)
        R.fig(savefig(fig, '04_tmq_t0_t3.png'), f'Ba thang con TMQ ở đầu buổi và sau 6 tuần (n = {len(base)}; thanh sai số = 95% KTC)')

    # ---------------- 8. Hao hụt mẫu
    R.h('8. Phân tích hao hụt mẫu')
    att_rows = []
    for lab, flag in (('Trả lời T2', 'co_t2'), ('Trả lời T3', 'co_t3')):
        stay, gone = t0[t0[flag] == 1], t0[t0[flag] == 0]
        for k, nm in (('t0_TOT', 'TMQ toàn thang (T0)'), ('t0_SE', 'Tự tin (T0)'), ('t0_KT', 'Kiến thức (T0)'), ('t1_YD', 'Ý định (T1)')):
            r = indep(stay[k], gone[k])
            att_rows.append({'Đợt': lab, 'Biến nền': nm, 'Ở lại n': r['n1'], 'Ở lại M': fnum(r['M1']), 'Rời đi n': r['n2'],
                             'Rời đi M': fnum(r['M2']), 'Hedges g': fnum(r.get('g')), 'p': fp(r.get('p'))})
    R.table('So sánh đặc điểm nền giữa người ở lại và người rời khỏi nghiên cứu', pd.DataFrame(att_rows), 'hao_hut',
            note='Nếu người ở lại có điểm nền cao hơn rõ rệt, kết quả T2/T3 có thể phóng đại mức vận dụng chung (thiên lệch tự chọn).')

    # ---------------- 9. Hoạt động workshop
    R.h('9. Dữ liệu từ các hoạt động trong workshop (mô tả)')
    hd = s[s.co_t0.eq(1) | s.co_t1.eq(1)]
    sh = hd['hd_hinh_1'].replace('', np.nan).dropna()
    rows9 = [{'Hoạt động': 'Hình khối chọn đầu tiên', 'Nội dung': SHAPES.get(k, k), 'n': int(v), '%': fpct(v / len(sh))}
             for k, v in sh.value_counts().items()]
    for i in range(1, 9):
        v = hd[f'hd_eisen_c{i}'].dropna()
        rows9.append({'Hoạt động': 'Eisenhower — % xếp đúng ô gợi ý', 'Nội dung': EISEN_NAME[f'c{i}'], 'n': len(v),
                      '%': fpct((v == EISEN_KEY[f'c{i}']).mean()) if len(v) else '—'})
    for k, nm in zip('SMART', ['Cụ thể', 'Đo được', 'Khả thi', 'Phù hợp', 'Có hạn']):
        v = hd['hd_muctieu_' + k].dropna()
        rows9.append({'Hoạt động': 'Mục tiêu — % tự đánh dấu đạt', 'Nội dung': nm, 'n': len(v), '%': fpct(v.mean()) if len(v) else '—'})
    for i, nm in enumerate(WHEEL, 1):
        v = hd[f'hd_banhxe{i}'].dropna()
        rows9.append({'Hoạt động': 'Bánh xe cuộc đời — điểm TB (1–10)', 'Nội dung': nm, 'n': len(v), '%': fnum(v.mean())})
    pk = hd['hd_camket'].dropna()
    for i, nm in enumerate(PLEDGES):
        rows9.append({'Hoạt động': 'Cam kết 7 ngày đã chọn', 'Nội dung': nm, 'n': int((pk == i).sum()), '%': fpct((pk == i).mean()) if len(pk) else '—'})
    for col, opts, nm in (('hd_cam', CAM, 'Cảm nhận thời gian mấy tuần đầu'), ('hd_dien_thoai', MAN, 'Thời gian dùng điện thoại hôm qua')):
        v = hd[col].dropna()
        for i, o in enumerate(opts):
            rows9.append({'Hoạt động': nm, 'Nội dung': o, 'n': int((v == i).sum()), '%': fpct((v == i).mean()) if len(v) else '—'})
    R.table('Kết quả các hoạt động tương tác (người đã đồng ý tham gia nghiên cứu)', pd.DataFrame(rows9), 'hoat_dong',
            note='Cột % của Bánh xe cuộc đời là điểm trung bình. Hình khối (Psycho-Geometrics) chỉ là công cụ khởi động, không dùng làm biến đo lường.')

    # ---------------- 10. Câu trả lời mở
    R.h('10. Câu trả lời mở')
    op = []
    for col, nm in (('hd_muctieu_noidung', 'Mục tiêu SMART tự viết'), ('hd_danhgia_noidung', 'Góp ý cuối buổi'),
                    ('t2_y_kien', 'Ý kiến sau 2 tuần'), ('t3_y_kien', 'Ý kiến sau 6 tuần')):
        for code, txt in s[['ma_an_danh', col]].itertuples(index=False):
            if isinstance(txt, str) and txt.strip():
                op.append({'Nguồn': nm, 'Mã': code, 'Nội dung': txt.strip(), 'Mã hóa (điền tay)': ''})
    opdf = pd.DataFrame(op, columns=['Nguồn', 'Mã', 'Nội dung', 'Mã hóa (điền tay)'])
    R.sheets['Tra_loi_mo'] = opdf
    R.p(f'Có {len(opdf)} câu trả lời mở, xuất sang trang “Tra_loi_mo” của tệp Excel để mã hóa theo chủ đề '
        '(gợi ý: hai người mã hóa độc lập, tính hệ số Cohen κ). Khi trích dẫn, bỏ mọi chi tiết có thể nhận ra người viết.')
    R.sheets['Du_lieu_da_cham'] = s
    return R


# ----------------------------------------------------------------------------- xuất tệp
def write_docx(R: Report, path: str, synthetic: bool, src: str, n: int) -> None:
    from docx import Document
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.LANDSCAPE
    sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    for m in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(sec, m, Cm(2))
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    for hs in ('Heading 1', 'Heading 2', 'Title'):
        doc.styles[hs].font.name = 'Times New Roman'
        doc.styles[hs].font.color.rgb = RGBColor(0x0F, 0x6B, 0x55)

    doc.add_heading('Báo cáo phân tích dữ liệu — Nghiên cứu kỹ năng quản lý thời gian của tân sinh viên K2026', level=0)
    if synthetic:
        p = doc.add_paragraph()
        r = p.add_run('DỮ LIỆU GIẢ LẬP — chỉ để kiểm tra bộ phân tích. Mọi con số dưới đây là số giả, KHÔNG dùng cho bài báo.')
        r.bold = True
        r.font.color.rgb = RGBColor(0xD2, 0x49, 0x2A)
    doc.add_paragraph(f'Tệp dữ liệu: {os.path.basename(src)} · {n} mã ẩn danh · Tạo lúc {dt.datetime.now():%d/%m/%Y %H:%M}')
    doc.add_heading('Ghi chú phương pháp phân tích', level=1)
    for t in [
        'TMQ (Britton & Tesser, 1991): 18 câu, thang 1–5; câu 8, 10, 11, 13, 16 chấm đảo (6 − x); điểm thang = trung bình các câu, điểm cao = quản lý thời gian tốt hơn.',
        'Kiến thức: 8 câu trắc nghiệm, mỗi câu đúng 1 điểm (0–8). Tự tin, phản hồi, ý định: trung bình các câu, thang 1–5. Vận dụng: trung bình 7 kỹ thuật, thang 0–3.',
        'So sánh ghép cặp: kiểm định t ghép cặp kèm d_z và 95% khoảng tin cậy (KTC); kiểm định Wilcoxon dấu-hạng để đối chiếu vì thang thứ bậc. '
        'So sánh hai nhóm độc lập: Mann–Whitney kèm Hedges g. Câu kiến thức: McNemar chính xác.',
        'Đa kiểm định: hiệu chỉnh Holm trong từng họ giả thuyết. Mức ý nghĩa α = 0,05 (hai phía). Chỉ người có dữ liệu ở cả hai thời điểm được đưa vào so sánh ghép cặp.',
        'Độ tin cậy: Cronbach α và McDonald ω (mô hình một nhân tố). Cấu trúc nhân tố TMQ: CFA bằng R/lavaan (tệp cfa_tmq.R).',
    ]:
        doc.add_paragraph(t, style='List Bullet')

    for b in R.blocks:
        if b[0] == 'h':
            doc.add_heading(b[1], level=1)
        elif b[0] == 'p':
            doc.add_paragraph(b[1])
        elif b[0] == 'img':
            doc.add_picture(b[1], width=Cm(22))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            cp = doc.add_paragraph()
            cr = cp.add_run(b[2])
            cr.italic = True
            cr.font.size = Pt(11)
        elif b[0] == 't':
            _, title, df, note = b
            tp = doc.add_paragraph()
            tr = tp.add_run(title)
            tr.bold = True
            tbl = doc.add_table(rows=1, cols=len(df.columns))
            tbl.style = 'Table Grid'
            for j, c in enumerate(df.columns):
                cell = tbl.rows[0].cells[j]
                cell.text = str(c)
                for pr in cell.paragraphs:
                    for rr in pr.runs:
                        rr.bold = True
                        rr.font.size = Pt(10)
            for row in df.itertuples(index=False):
                cells = tbl.add_row().cells
                for j, v in enumerate(row):
                    cells[j].text = '' if v is None or (isinstance(v, float) and math.isnan(v)) else str(v)
                    for pr in cells[j].paragraphs:
                        for rr in pr.runs:
                            rr.font.size = Pt(10)
            if note:
                npar = doc.add_paragraph()
                nr = npar.add_run('Ghi chú. ' + note)
                nr.italic = True
                nr.font.size = Pt(10)
    doc.save(path)


CFA_R = r'''# CFA 3 nhân tố cho TMQ (Britton & Tesser, 1991) — dữ liệu T0 đã đảo chiều trong du_lieu_da_cham.csv
# Cài một lần: install.packages(c("lavaan", "semTools"))
library(lavaan)
d <- read.csv("du_lieu_da_cham.csv", fileEncoding = "UTF-8-BOM")
d <- d[d$co_t0 == 1, ]
items <- sprintf("t0_q%02d", 1:18)
mo_hinh <- '
  SRP =~ t0_q01 + t0_q02 + t0_q03 + t0_q04 + t0_q05 + t0_q06 + t0_q07
  TA  =~ t0_q08 + t0_q09 + t0_q10 + t0_q11 + t0_q12 + t0_q13
  LRP =~ t0_q14 + t0_q15 + t0_q16 + t0_q17 + t0_q18
'
fit <- cfa(mo_hinh, data = d, ordered = items, estimator = "WLSMV")
summary(fit, fit.measures = TRUE, standardized = TRUE)
fitMeasures(fit, c("chisq.scaled", "df", "pvalue.scaled", "cfi.scaled", "tli.scaled", "rmsea.scaled", "srmr"))
# Ngưỡng tham khảo: CFI/TLI ≥ 0,90 (chấp nhận) / ≥ 0,95 (tốt); RMSEA ≤ 0,08; SRMR ≤ 0,08.
semTools::reliability(fit)   # ω cho từng nhân tố
modindices(fit, sort. = TRUE, maximum.number = 10)

# Bất biến đo lường T0–T3 (khi đủ cỡ mẫu ghép cặp, khoảng ≥ 150): dựng mô hình dọc với
# hệ số tải bằng nhau giữa hai thời điểm và so sánh ΔCFI (< 0,01).
'''


def main():
    global DEC
    ap = argparse.ArgumentParser(description='Phân tích dữ liệu nghiên cứu Workshop K2026')
    ap.add_argument('csv', nargs='?', help='tệp CSV tải từ bảng diễn giả')
    ap.add_argument('--ra', help='thư mục kết quả (mặc định: ket_qua_<ngày giờ> cạnh tệp CSV)')
    ap.add_argument('--thap-phan', default=',', choices=[',', '.'], help='dấu thập phân trong báo cáo (mặc định: dấu phẩy)')
    ap.add_argument('--loai-mot-kieu', action='store_true', help='loại người trả lời TMQ một kiểu (18 câu cùng mức) — phân tích độ nhạy')
    ap.add_argument('--gia-lap', metavar='TEP_RA', help='tạo dữ liệu giả lập để chạy thử, rồi dừng')
    ap.add_argument('--n', type=int, default=230, help='số người có T0 khi giả lập (mặc định 230)')
    a = ap.parse_args()
    if a.gia_lap:
        simulate(a.gia_lap, a.n)
        return
    if not a.csv:
        ap.error('cần đường dẫn tệp CSV (hoặc dùng --gia-lap)')
    DEC = a.thap_phan
    df = load(a.csv)
    s = score(df)
    synthetic = len(df) > 0 and df['ma_an_danh'].astype(str).str.startswith('ZZ').all()
    out = a.ra or os.path.join(os.path.dirname(os.path.abspath(a.csv)), 'ket_qua_' + dt.datetime.now().strftime('%Y%m%d_%H%M'))
    os.makedirs(out, exist_ok=True)
    R = analyse(df, s, out, synthetic, a.csv, a.loai_mot_kieu)
    s.to_csv(os.path.join(out, 'du_lieu_da_cham.csv'), index=False, encoding='utf-8-sig')
    with open(os.path.join(out, 'cfa_tmq.R'), 'w', encoding='utf-8') as f:
        f.write(CFA_R)
    with pd.ExcelWriter(os.path.join(out, 'ket_qua_phan_tich.xlsx'), engine='openpyxl') as xw:
        info = pd.DataFrame({'Thông tin': ['Tệp dữ liệu', 'Số mã ẩn danh', 'Thời điểm phân tích', 'Loại dữ liệu'],
                             'Giá trị': [os.path.basename(a.csv), len(df), dt.datetime.now().strftime('%d/%m/%Y %H:%M'),
                                         'DỮ LIỆU GIẢ LẬP — không dùng để công bố' if synthetic else 'Dữ liệu thật']})
        info.to_excel(xw, sheet_name='Thong_tin', index=False)
        for name, d in R.sheets.items():
            d.to_excel(xw, sheet_name=name[:31], index=False)
        for ws in xw.book.worksheets:
            for col in ws.columns:
                width = max(len(str(c.value)) if c.value is not None else 0 for c in col[:60])
                ws.column_dimensions[col[0].column_letter].width = min(max(10, width + 2), 60)
            ws.freeze_panes = 'A2'
    write_docx(R, os.path.join(out, 'bao_cao_phan_tich.docx'), synthetic, a.csv, len(df))
    print(('[GIẢ LẬP] ' if synthetic else '') + f'Xong. Kết quả trong: {out}')


if __name__ == '__main__':
    main()
